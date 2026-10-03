#!/usr/bin/env python3
"""Precision, recall and error attribution from the clinician's evaluation.

The reference clinician graded every model code. Verdicts are read from each
model's evaluation workbook and counted as:

    TP         correct and clinically central
    FP_soft    valid for the patient but not central to the admission
    FP_hard    not applicable to the patient (hallucination)
    FN_coding  concept found in the text but given the wrong or no code
    FN_miss    relevant concept not extracted at all
    TN         ruled-out finding correctly not coded

Validity credits FP_soft; clinical alignment does not. Definitions are in
Supplementary Material B. Writes adjudication.csv and adjudication_counts.csv.

    python -m analysis.adjudication --study config/study.toml
"""

import argparse
from collections import Counter

import pandas as pd
from openpyxl import load_workbook

from analysis.data import load_study

VERDICTS = {"TP", "TN", "FP_soft", "FP_hard", "FN_miss", "FN_coding", "FN_other", "AMBIG"}
MISS_PHRASES = ("does not code", "relevant here", "would still be relevant", "no diagnosis",
                "diagnosis not", "not included", "no procedure", "procedure not")


def classify(label):
    """Map the clinician's free-text verdict to a category.

    Args:
        label: Cell value from the verdict column.

    Returns:
        A category in `VERDICTS`, "UNMAPPED:<text>", or None for blanks and
        notes that are not verdicts.
    """
    if label is None:
        return None
    text = str(label).strip().lower()
    if not text:
        return None
    if text.startswith("tf"):
        return "AMBIG"
    if text.startswith("tn"):
        return "TN"
    if text.startswith("tp"):
        return "TP"
    if "fn" in text:
        if "correct" in text and ("extract" in text or "text" in text):
            return "FN_coding"
        if any(phrase in text for phrase in MISS_PHRASES):
            return "FN_miss"
        if "incorrect code" in text or "no code" in text:
            return "FN_coding"
        return "FN_other"
    if "fp" in text:
        return "FP_soft" if "additional not relevant" in text else "FP_hard"
    if not any(t in text for t in ("tp", "fp", "fn", "tn")):
        return None
    return "UNMAPPED:" + str(label)[:40]


def _verdict_hits(sheet, column):
    return sum(1 for r in range(2, sheet.max_row + 1)
               if str(classify(sheet.cell(r, column).value)).split(":")[0] in VERDICTS)


def load_verdicts(path):
    """Read verdicts from an evaluation workbook.

    The verdict column is the one, across non-template sheets, holding the
    most recognisable verdicts.

    Args:
        path: Evaluation workbook.

    Returns:
        List of (category, qualifier, primary_secondary) per graded row.
    """
    workbook = load_workbook(path, data_only=True)
    best = None
    for name in workbook.sheetnames:
        if "template" in name.lower():
            continue
        sheet = workbook[name]
        for column in range(1, sheet.max_column + 1):
            hits = _verdict_hits(sheet, column)
            if best is None or hits > best[0]:
                best = (hits, name, column)
    if best is None or best[0] == 0:
        raise ValueError(f"No verdict column found in {path}")
    _, name, verdict_col = best
    sheet = workbook[name]
    header = {str(sheet.cell(1, c).value).strip().lower(): c
              for c in range(1, sheet.max_column + 1) if sheet.cell(1, c).value}
    qualifier_col = next((c for k, c in header.items() if "qualifier" in k), None)
    position_col = next((c for k, c in header.items() if "primary" in k), None)

    rows = []
    for r in range(2, sheet.max_row + 1):
        category = classify(sheet.cell(r, verdict_col).value)
        if category is None:
            continue
        qualifier = _cell(sheet, r, qualifier_col)
        position = _cell(sheet, r, position_col)
        rows.append((category, qualifier, position))
    return rows


def _cell(sheet, row, column):
    return str(sheet.cell(row, column).value).strip().lower() if column else ""


def _prf(tp, fp, fn):
    p = tp / (tp + fp) if tp + fp else float("nan")
    r = tp / (tp + fn) if tp + fn else float("nan")
    f = 2 * p * r / (p + r) if (p and r and p + r) else float("nan")
    return p, r, f


def evaluate(rows):
    """Compute the evaluation metrics from graded rows.

    Args:
        rows: Output of `load_verdicts`.

    Returns:
        Dict of counts and metrics.
    """
    counts = Counter(category for category, _, _ in rows)
    tp, fp_soft, fp_hard = counts["TP"], counts["FP_soft"], counts["FP_hard"]
    fn_coding = counts["FN_coding"]
    fn_miss = counts["FN_miss"] + counts["FN_other"]
    fn = fn_coding + fn_miss
    alignment = _prf(tp, fp_soft + fp_hard, fn)
    validity = _prf(tp + fp_soft, fp_hard, fn)
    produced = tp + fp_soft + fp_hard
    found = tp + fn_coding

    def nan_div(a, b):
        return a / b if b else float("nan")

    return {
        "n": len(rows), "tp": tp, "fp_soft": fp_soft, "fp_hard": fp_hard,
        "fn_coding": fn_coding, "fn_miss": fn_miss, "tn": counts["TN"],
        "validity_precision": validity[0], "validity_recall": validity[1],
        "validity_f1": validity[2],
        "alignment_precision": alignment[0], "alignment_recall": alignment[1],
        "alignment_f1": alignment[2],
        "extraction_recall": nan_div(found, found + fn_miss),
        "coding_accuracy": nan_div(tp, found),
        "hallucination_rate": nan_div(fp_hard, produced),
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__,
                                     formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--study", default="config/study.toml")
    args = parser.parse_args()

    study = load_study(args.study)
    results, counts = [], []
    for name, paths in study["models"].items():
        rows = load_verdicts(paths["evaluation"])
        results.append({"model": name, **evaluate(rows)})
        for category, n in sorted(Counter(c for c, _, _ in rows).items()):
            counts.append({"model": name, "category": category, "n": n})

    table = pd.DataFrame(results)
    table.to_csv(study["output_dir"] / "adjudication.csv", index=False)
    pd.DataFrame(counts).to_csv(study["output_dir"] / "adjudication_counts.csv", index=False)
    pd.set_option("display.width", 200)
    print(table.round(2).to_string(index=False))


if __name__ == "__main__":
    main()
